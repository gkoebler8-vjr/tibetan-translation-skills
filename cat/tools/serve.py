#!/usr/bin/env python3
"""
serve.py -- Tiger CAT: the local app server. One folder of projects; each project is a text, its brief,
the skill's run directory, the translator's edits and the exports.

  python3 cat/tools/serve.py [--projects cat/projects] [--port 8765]     (from the repository root; or tools/serve.py from cat/)

A project folder:  projects/<slug>/project.json  (title, brief, model, effort, page size)
                   projects/<slug>/run/           the skill's working directory: <slug>.units.md, CLAUDE.md (brief),
                                                  final.md, <slug>.glossary.tsv, <slug>.sources.tsv, dm_*.txt, runlog.md
                   projects/<slug>/data.json      built from run/ by tools/build_data.py (rebuilt when run/ changes)
                   projects/<slug>/state.json     the translator's edits        projects/<slug>/export/   exports
                   projects/<slug>/jobs/          logs of translation runs

Translation runs are executed by the user's own Claude Code binary in headless mode (claude -p) inside run/,
with the tibetan-translate skill, under the user's own login. The server never calls a model or Dharmamitra
itself. Binds to 127.0.0.1 only.
"""
import argparse, datetime, glob, http.server, json, os, re, shutil, subprocess, sys, threading, time, traceback, urllib.parse

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
APP = os.path.join(ROOT, 'app')
REPO = os.path.dirname(ROOT)                      # the skills repository: cat/ sits inside it
INSTALLED_SKILL = os.path.expanduser('~/.claude/skills/tibetan-translate')   # the runtime copy headless Claude Code loads
REPO_SKILL = os.path.join(REPO, 'skills', 'tibetan-translate')              # the source of truth
# export_docx.py comes from the installed skill; fall back to the repo copy on a fresh checkout
SKILL_TOOLS = os.path.join(INSTALLED_SKILL, 'tools') if os.path.isdir(os.path.join(INSTALLED_SKILL, 'tools')) else os.path.join(REPO_SKILL, 'tools')
PROJECTS = os.path.join(ROOT, 'projects')
sys.path.insert(0, HERE)
import tibseg

def load(path, default):
    try: return json.load(open(path, encoding='utf-8'))
    except Exception: return default
def dump(path, obj): json.dump(obj, open(path, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
def now(): return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

# ------------------------------------------------------------------ interpreters and the Claude binary
def build_python():
    """An interpreter with pyewts (script <-> Wylie). Prefer our own venv, then the skill's, then this one."""
    for c in [os.path.expanduser('~/.venvs/vcat/bin/python3'), os.path.expanduser('~/.venvs/tib/bin/python3')]:
        if os.path.exists(c): return c
    return sys.executable

def claude_binary(project=None):
    if project and project.get('claude_path') and os.path.exists(project['claude_path']): return project['claude_path']
    env = os.environ.get('VCAT_CLAUDE')
    if env and os.path.exists(env): return env
    for c in [shutil.which('claude'), os.path.expanduser('~/.claude/local/claude'), os.path.expanduser('~/.local/bin/claude')]:
        if c and os.path.exists(c): return c
    cands = sorted(glob.glob(os.path.expanduser('~/Library/Application Support/Claude/claude-code/*/*/claude.app/Contents/MacOS/claude')))
    return cands[-1] if cands else ''

# ------------------------------------------------------------------ projects
def slugify(name):
    s = re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')
    return s[:40] or 'text'
def pdir(slug): return os.path.join(PROJECTS, slug)
def pjson(slug): return load(os.path.join(pdir(slug), 'project.json'), None)
def run_dir(slug): return os.path.join(pdir(slug), 'run')
def units_path(slug): return os.path.join(run_dir(slug), slug + '.units.md')

BRIEF_DEFAULTS = {'audience': 'seasoned practitioner', 'purpose': 'study aid', 'style': '', 'notes': 'reader', 'context': '', 'prior': 'consult', 'language': 'en'}
LANG_NAMES = {'en': 'English', 'de': 'German'}
AUDIENCES = ['academic', 'new practitioner', 'seasoned practitioner', 'hybrid']
PURPOSES = ['publication', 'study aid', 'practice/recitation', 'working crib']
NOTES = ['scholarly', 'reader', 'minimal', 'none']

def list_projects():
    out = []
    for d in sorted(glob.glob(os.path.join(PROJECTS, '*'))):
        p = load(os.path.join(d, 'project.json'), None)
        if not p: continue
        data = load(os.path.join(d, 'data.json'), {})
        n = len(data.get('units', [])); pend = data.get('pending', 0)
        out.append({'slug': os.path.basename(d), 'title': p.get('title', os.path.basename(d)), 'units': n, 'pending': pend, 'model': p.get('model'), 'effort': p.get('effort'), 'updated': p.get('updated', '')})
    out.sort(key=lambda x: x['updated'], reverse=True)
    return out

def write_brief(slug, p):
    b = p.get('brief', {})
    lines = ['# %s' % p.get('title', slug), '',
             'This folder is a translation project of Tiger CAT. The brief below answers Pass 0 of the tibetan-translate skill; do not ask it again.', '',
             '## Brief', '',
             '- Target language: %s%s' % (LANG_NAMES.get(b.get('language', 'en'), 'English'), '' if b.get('language', 'en') == 'en' else ' (render the TEXT, the footnotes and the fixed short titles in this language; keep the editor notes in English; bind the glossary in this language; apply reference/german.md where it exists, otherwise the English rules mutatis mutandis)'),
             '- Audience: %s' % b.get('audience', ''), '- Purpose: %s' % b.get('purpose', ''),
             '- House style: %s' % (b.get('style') or 'per reference/modes.md for the audience'),
             '- Notes policy: %s footnotes; sources in the register (%s.sources.tsv)' % (b.get('notes', 'reader'), slug),
             '- Source context: %s' % (b.get('context') or 'infer from the text and state the assumption'),
             '- Existing translations: %s' % ({'ignore': 'ignore them (do not read or cite)', 'consult': 'consult where found and cite them (dm.py cite)', 'adapt': 'may adapt under the permission recorded here'}.get(b.get('prior', 'consult'))),
             '- Glossary: %s.glossary.tsv binds; append every term decision and TITLE row to it.' % slug,
             '- Page file: %s.units.md (one Uxx line per unit); output: final.md (append unit blocks in order, never rewrite earlier ones); runlog.md; save every dm.py output as dm_<command>_<unit>.txt here.' % slug, '']
    open(os.path.join(run_dir(slug), 'CLAUDE.md'), 'w', encoding='utf-8').write('\n'.join(lines))

def write_units(slug, lines, title):
    hdr = ['# %s' % title, '', 'Units (one per line; the parenthesis after the id is a hint, not part of the text).', '']
    open(units_path(slug), 'w', encoding='utf-8').write('\n'.join(hdr + lines) + '\n')

def to_script_if_wylie(text):
    if re.search(r'[ༀ-࿿]', text): return text
    code = "import sys,pyewts; c=pyewts.pyewts(); print(c.toUnicode(sys.stdin.read()))"
    try:
        r = subprocess.run([build_python(), '-I', '-c', code], input=text, capture_output=True, text=True, timeout=60)
        if r.returncode == 0 and r.stdout.strip(): return r.stdout
    except Exception: pass
    return text

def rebuild(slug):
    p = pjson(slug)
    if not p: return False, 'no project'
    cmd = [build_python(), os.path.join(HERE, 'build_data.py'), run_dir(slug), '--units', units_path(slug), '--work', slug, '-o', os.path.join(pdir(slug), 'data.json'),
           '--title', p.get('title', slug), '--brief', ' · '.join(x for x in [p.get('brief', {}).get('audience', ''), p.get('brief', {}).get('purpose', ''), (p.get('brief', {}).get('style') or '')[:60]] if x)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    return r.returncode == 0, (r.stdout + r.stderr).strip()[-600:]

def create_project(body):
    title = (body.get('title') or '').strip() or 'Untitled text'
    slug = slugify(title); base = slug; k = 2
    while os.path.exists(pdir(slug)): slug = '%s-%d' % (base, k); k += 1
    os.makedirs(run_dir(slug)); os.makedirs(os.path.join(pdir(slug), 'export'), exist_ok=True); os.makedirs(os.path.join(pdir(slug), 'jobs'), exist_ok=True)
    text = to_script_if_wylie(body.get('text') or '')
    max_syl = max(20, min(120, int(body.get('max_syl') or 45)))
    units = tibseg.segment(text, max_syl)
    p = {'title': title, 'brief': {**BRIEF_DEFAULTS, **(body.get('brief') or {})}, 'model': body.get('model') or 'opus', 'effort': body.get('effort') or 'xhigh',
         'page_size': int(body.get('page_size') or 10), 'max_syl': max_syl, 'created': now(), 'updated': now(), 'source_text': text}
    dump(os.path.join(pdir(slug), 'project.json'), p)
    write_units(slug, tibseg.to_lines(units), title); write_brief(slug, p)
    open(os.path.join(run_dir(slug), slug + '.glossary.tsv'), 'w', encoding='utf-8').write('wylie\t%s\tnote\n' % LANG_NAMES.get(p['brief'].get('language', 'en'), 'English'))
    if body.get('glossary'): import_glossary(slug, body['glossary'], rebuild_after=False)
    ok, log = rebuild(slug)
    return {'ok': ok, 'slug': slug, 'units': len(units), 'log': log}

def update_project(slug, body):
    p = pjson(slug)
    if not p: return {'ok': False, 'error': 'no project'}
    for k in ('title', 'model', 'effort', 'claude_path'):
        if k in body: p[k] = body[k]
    if 'page_size' in body: p['page_size'] = int(body['page_size'] or 10)
    if 'max_syl' in body: p['max_syl'] = max(20, min(120, int(body['max_syl'] or 45)))
    if 'brief' in body: p['brief'] = {**p.get('brief', {}), **body['brief']}
    p['updated'] = now(); dump(os.path.join(pdir(slug), 'project.json'), p); write_brief(slug, p)
    if 'units' in body:
        lines = []
        for i, l in enumerate([x for x in body['units'].splitlines() if x.strip()]):
            m = re.match(r'^(U\d+[a-z]?)\s*(\([^)]*\))?\s*(.*)$', l.strip())
            if m and m.group(3): lines.append('%s %s %s' % (m.group(1), m.group(2) or '(prose)', m.group(3)))
            else: lines.append('U%02d (prose) %s' % (i + 1, l.strip()))
        write_units(slug, lines, p['title'])
    ok, log = rebuild(slug)
    return {'ok': ok, 'log': log}

def resegment_project(slug, body):
    p = pjson(slug)
    if not p: return {'ok': False, 'error': 'no project'}
    text = body.get('text') or p.get('source_text') or ''
    if not text.strip(): return {'ok': False, 'error': 'no source text stored for this project'}
    max_syl = max(20, min(120, int(body.get('max_syl') or p.get('max_syl') or 45)))
    units = tibseg.segment(to_script_if_wylie(text), max_syl)
    p['source_text'] = text; p['max_syl'] = max_syl; p['updated'] = now(); dump(os.path.join(pdir(slug), 'project.json'), p)
    write_units(slug, tibseg.to_lines(units), p['title'])
    ok, log = rebuild(slug)
    return {'ok': ok, 'units': len(units), 'log': log}

def delete_project(slug):
    d = pdir(slug)
    if os.path.exists(d): shutil.rmtree(d)
    return {'ok': True}

# ------------------------------------------------------------------ glossary import
def to_wylie_many(texts):
    if not any(re.search(r'[ༀ-࿿]', t) for t in texts): return texts
    code = ("import sys,json,re\nimport pyewts\nc=pyewts.pyewts()\nxs=json.load(sys.stdin)\nout=[]\nfor t in xs:\n  if re.search(r'[\\u0F00-\\u0FFF]',t):\n    w=c.toWylie(t).replace('_',' ')\n    w=re.sub(r'/\\s*/','//',w)\n    out.append(re.sub(r'\\s+',' ',w).strip(' /'))\n  else: out.append(t)\njson.dump(out,sys.stdout,ensure_ascii=False)")
    try:
        r = subprocess.run([build_python(), '-I', '-c', code], input=json.dumps(texts, ensure_ascii=False), capture_output=True, text=True, timeout=60)
        if r.returncode == 0: return json.loads(r.stdout)
    except Exception: pass
    return texts

def import_glossary(slug, text, rebuild_after=True):
    gpath = os.path.join(run_dir(slug), slug + '.glossary.tsv')
    rows = []
    for line in (text or '').splitlines():
        line = line.strip('﻿').rstrip('\n')
        if not line.strip() or line.lstrip().startswith('#'): continue
        cols = line.split('\t') if '\t' in line else re.split(r'\s*[;,]\s*', line, maxsplit=2) if re.search(r'[;,]', line) else [line]
        cols = [c.strip().strip('"') for c in cols]
        if not rows and len(cols) > 1 and cols[0].lower() in ('wylie', 'tibetan', 'term', 'bo', 'source') and cols[1].lower() in ('english', 'en', 'target', 'translation', 'gloss'): continue
        if len(cols) < 2 or not cols[0] or not cols[1]: continue
        rows.append(cols[:3] + [''] * (3 - len(cols[:3])))
    if not rows: return {'ok': False, 'error': 'no rows found (expected: tibetan<TAB>english<TAB>note)'}
    wylies = to_wylie_many([r[0] for r in rows])
    existing = set()
    if os.path.exists(gpath):
        for line in open(gpath, encoding='utf-8'):
            c = line.rstrip('\n').split('\t')
            if c and c[0] not in ('wylie', 'TITLE'): existing.add(c[0].strip().lower())
    else: open(gpath, 'w', encoding='utf-8').write('wylie\tEnglish\tnote\n')
    added = skipped = 0
    with open(gpath, 'a', encoding='utf-8') as f:
        for (tib, en, note), wy in zip(rows, wylies):
            if wy.strip().lower() in existing: skipped += 1; continue
            f.write('%s\t%s\t%s\n' % (wy.strip(), en.strip(), (note.strip() + ' ' if note.strip() else '') + '[imported %s]' % datetime.date.today().isoformat()))
            existing.add(wy.strip().lower()); added += 1
    res = {'ok': True, 'added': added, 'skipped': skipped, 'file': gpath}
    if rebuild_after: ok, log = rebuild(slug); res['ok'] = ok; res['log'] = log
    return res

# ------------------------------------------------------------------ refresh
def run_stamp(slug):
    files = [os.path.join(run_dir(slug), f) for f in ('final.md', slug + '.glossary.tsv', slug + '.sources.tsv', slug + '.units.md')]
    return max([os.path.getmtime(f) for f in files if os.path.exists(f)] or [0])

def api_refresh(slug):
    dp = os.path.join(pdir(slug), 'data.json')
    stamp = run_stamp(slug); built = os.path.getmtime(dp) if os.path.exists(dp) else 0
    if stamp > built + 1:
        ok, log = rebuild(slug); return {'rebuilt': ok, 'log': log, 'stamp': stamp, 'job': job_status(slug)}
    return {'rebuilt': False, 'stamp': stamp, 'job': job_status(slug)}

# ------------------------------------------------------------------ translation runs through the user's own Claude Code
JOBS = {}   # slug -> dict
JOBLOCK = threading.Lock()

def pages_of(slug, p):
    data = load(os.path.join(pdir(slug), 'data.json'), {})
    pending = [u['id'] for u in data.get('units', []) if u.get('pending')]
    size = max(3, min(12, int(p.get('page_size') or 10)))
    return [pending[i:i + size] for i in range(0, len(pending), size)]

def page_prompt(slug, p, ids):
    b = p.get('brief', {})
    return '\n'.join([
        'Translate units %s–%s of the page file %s with the tibetan-translate skill (pipeline v2), end to end, exactly as the skill prescribes.' % (ids[0], ids[-1], units_path(slug)),
        'Invoke the skill with the Skill tool first, then read ~/.claude/skills/tibetan-translate/SKILL.md from disk: the on-disk text is authoritative. Follow it, including the references it tells you to read and when.',
        'You are running autonomously inside Tiger CAT: do not ask questions; take the brief in CLAUDE.md of this folder as given (Pass 0 is answered); state the model/effort line as the skill asks and proceed.',
        'Run directory (all work files go here): %s' % run_dir(slug), 'Work name for the files: %s' % slug,
        'BRIEF: target language %s · audience %s · purpose %s · house style: %s · notes policy: %s · source context: %s · existing translations: %s.' % (LANG_NAMES.get(b.get('language', 'en'), 'English'), b.get('audience'), b.get('purpose'), b.get('style') or 'per modes.md', b.get('notes'), b.get('context') or 'infer and state', b.get('prior')),
        'Treat these units as ONE PAGE: your own reading and draft of all of them with the construal sketch in %s.construal.md and the glossary in %s.glossary.tsv (it binds; append); then the grounding pass; then the style pass; then the in-context check; then notes and footnotes with the confidence grade.' % (slug, slug),
        'Append the finished unit blocks (### Uxx / HEADER / TEXT / FOOTNOTES / NOTES) to %s/final.md, creating it with a title line if it does not exist; never rewrite earlier units. Keep runlog.md; save every dm.py output as dm_<command>_<unit>.txt here. Register rows go to %s.sources.tsv.' % (run_dir(slug), slug),
        'Reply with only: the check line and the grounding summary in two lines. Do not paste the translation into the reply.'])

def parse_reset(text):
    """A reset time from a usage-limit message, if one is printed; else None."""
    m = re.search(r'resets?\s*(?:at\s*)?(\d{1,2})(?::(\d{2}))?\s*(am|pm)?', text, re.I)
    if not m: return None
    h = int(m.group(1)); mi = int(m.group(2) or 0); ap = (m.group(3) or '').lower()
    if ap == 'pm' and h < 12: h += 12
    if ap == 'am' and h == 12: h = 0
    t = datetime.datetime.now().replace(hour=h % 24, minute=mi, second=0, microsecond=0)
    if t < datetime.datetime.now(): t += datetime.timedelta(days=1)
    return t

def parse_reset_at(text, base):
    m = re.search(r'resets?\s*(?:at\s*)?(\d{1,2})(?::(\d{2}))?\s*(am|pm)?', text, re.I)
    if not m: return None
    h = int(m.group(1)); mi = int(m.group(2) or 0); ap = (m.group(3) or '').lower()
    if ap == 'pm' and h < 12: h += 12
    if ap == 'am' and h == 12: h = 0
    t = base.replace(hour=h % 24, minute=mi, second=0, microsecond=0)
    if t < base: t += datetime.timedelta(days=1)
    return t

def usage_entries():
    """One entry per headless page run, from the job logs: when it started, what it used, whether it hit a limit."""
    out = []
    for log in glob.glob(os.path.join(PROJECTS, '*', 'jobs', '*.log')):
        ts = datetime.datetime.fromtimestamp(os.path.getmtime(log))
        try: lines = open(log, encoding='utf-8', errors='replace').read().splitlines()
        except Exception: continue
        for line in lines:
            if line.startswith('====='):
                m = re.match(r'===== (\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})', line)
                if m: ts = datetime.datetime.strptime(m.group(1), '%Y-%m-%d %H:%M:%S')
                continue
            if '"type":"result"' not in line: continue
            try: j = json.loads(line)
            except ValueError: continue
            if j.get('type') != 'result': continue
            u = j.get('usage') or {}; res = str(j.get('result') or '')
            lim = bool(j.get('is_error')) and bool(re.search(r'limit', res, re.I))
            out.append({'t': ts, 'in': u.get('input_tokens') or 0, 'out': u.get('output_tokens') or 0, 'cr': u.get('cache_read_input_tokens') or 0, 'cw': u.get('cache_creation_input_tokens') or 0,
                        'cost': j.get('total_cost_usd') or 0, 'limit': lim, 'weekly': lim and bool(re.search(r'week', res, re.I)), 'reset': parse_reset_at(res, ts) if lim else None, 'ok': not j.get('is_error')})
    out.sort(key=lambda e: e['t']); return out

def api_usage():
    es = usage_entries(); now = datetime.datetime.now()
    tok = lambda e: e['out'] + e['in'] + e['cw']
    sess_hits = [e for e in es if e['limit'] and not e['weekly']]; week_hits = [e for e in es if e['weekly']]
    last = sess_hits[-1] if sess_hits else None
    exhausted = bool(last and last['reset'] and now < last['reset'])
    if last and last['reset'] and now >= last['reset']: start = last['reset']
    elif exhausted: start = last['t'] - datetime.timedelta(hours=5)
    else: start = now - datetime.timedelta(hours=5)
    sess = [e for e in es if e['t'] >= start and (not exhausted or e['t'] <= last['t'])]
    ceiling = None
    if last:
        before = [e for e in es if last['t'] - datetime.timedelta(hours=5) < e['t'] <= last['t']]
        ceiling = sum(tok(e) for e in before) or None
    sess_tokens = sum(tok(e) for e in sess)
    if ceiling and not exhausted and sess_tokens > ceiling * 1.1: ceiling = None   # the learned ceiling is stale: this window already went past it
    wk = [e for e in es if e['t'] >= now - datetime.timedelta(days=7)]
    wceil = None
    if week_hits:
        wl = week_hits[-1]; wceil = sum(tok(e) for e in es if wl['t'] - datetime.timedelta(days=7) < e['t'] <= wl['t']) or None
    return {'session': {'tokens': sum(tok(e) for e in sess), 'cache_read': sum(e['cr'] for e in sess), 'pages': sum(1 for e in sess if e['ok']), 'ceiling': ceiling, 'exhausted': exhausted,
                        'reset': (last['reset'].strftime('%H:%M') if last and last['reset'] and now < last['reset'] else None)},
            'week': {'tokens': sum(tok(e) for e in wk), 'pages': sum(1 for e in wk if e['ok']), 'cost': sum(e['cost'] for e in wk), 'ceiling': wceil},
            'runs': len(es)}

STAGES = ['starting', 'reading the skill', 'reading and drafting', 'grounding', 'second opinion', 'style pass', 'checking', 'writing the page']

def stage_of(ev_name, ev_input):
    """Which pipeline stage a tool call belongs to; None when it says nothing about the stage."""
    cmd = str((ev_input or {}).get('command') or ''); fp = str((ev_input or {}).get('file_path') or ''); skill = str((ev_input or {}).get('skill') or '')
    if ev_name == 'Skill' or (ev_name == 'Read' and '/skills/' in fp and 'SKILL.md' in fp): return 1
    if ev_name in ('Write', 'Edit') and re.search(r'\.(construal|draft\d*)\.md$', fp): return 2
    if 'dm.py' in cmd:
        if re.search(r'dm\.py\s+translate', cmd): return 4
        if re.search(r'dm\.py\s+(identify|gloss|explore|segment|parallels|search|meta|cite)', cmd): return 3
    if 'tibdict.py' in cmd: return 2
    if ev_name in ('Write', 'Edit') and re.search(r'pass3|style', fp): return 5
    if 'beats.py' in cmd: return 5
    if ev_name == 'Read' and fp.endswith('check.md'): return 6
    if ev_name in ('Write', 'Edit') and fp.endswith('final.md'): return 7
    if ev_name == 'Bash' and re.search(r'>>\s*\S*final\.md', cmd): return 7
    return None

def page_times_path(): return os.path.join(PROJECTS, 'page-times.json')
def record_page_time(slug, ids, seconds, ok):
    rows = load(page_times_path(), [])
    rows.append({'slug': slug, 'ids': [ids[0], ids[-1]], 'n': len(ids), 'seconds': int(seconds), 'ok': ok, 't': now()})
    dump(page_times_path(), rows[-200:])
def typical_page_seconds():
    rows = [r['seconds'] for r in load(page_times_path(), []) if r.get('ok') and r.get('seconds', 0) > 120]
    if not rows: return None
    rows.sort(); return rows[len(rows) // 2]

def job_status(slug):
    j = JOBS.get(slug)
    if not j: return {'state': 'idle'}
    out = {k: j.get(k) for k in ('state', 'page', 'pages', 'ids', 'started', 'message', 'wait_until', 'log', 'last', 'needs_login', 'stage', 'grounding_calls')}
    out['stages'] = STAGES; out['typical_s'] = typical_page_seconds()
    ps = j.get('page_started'); out['elapsed_s'] = int((datetime.datetime.now() - ps).total_seconds()) if ps and j.get('state') == 'running' else None
    return out

def run_job(slug, scope):
    p = pjson(slug); claude = claude_binary(p)
    j = JOBS[slug]
    if not claude:
        j.update(state='error', message='Claude Code was not found. Install the claude command line (https://claude.ai/code) or set its path in the project settings.'); return
    if not claude_auth_status(slug).get('logged_in'):
        j.update(state='error', message='Claude Code is not signed in yet. Use "Sign in to Claude Code" in the Translate menu (one time), then start the run again.', needs_login=True); return
    pages = pages_of(slug, p)
    if not pages: j.update(state='done', message='Nothing left to translate.'); return
    if scope == 'next': pages = pages[:1]
    j.update(pages=len(pages))
    log_path = os.path.join(pdir(slug), 'jobs', datetime.datetime.now().strftime('%Y%m%d-%H%M%S') + '.log'); j['log'] = log_path
    env = dict(os.environ); env.pop('CLAUDECODE', None); env.pop('CLAUDE_CODE_ENTRYPOINT', None)
    for n, ids in enumerate(pages, 1):
        if j.get('stop'): j.update(state='stopped', message='Stopped before page %d.' % n); return
        attempts = 0
        while True:
            attempts += 1
            j.update(state='running', page=n, ids=[ids[0], ids[-1]], message='Translating %s–%s (page %d of %d)%s' % (ids[0], ids[-1], n, len(pages), ', attempt %d' % attempts if attempts > 1 else ''), last='', stage=0, grounding_calls=0, page_started=datetime.datetime.now())
            cmd = [claude, '-p', page_prompt(slug, p, ids), '--model', p.get('model') or 'opus', '--effort', p.get('effort') or 'xhigh',
                   '--permission-mode', 'bypassPermissions', '--add-dir', os.path.expanduser('~/.claude/skills'), '--output-format', 'stream-json', '--verbose']
            with open(log_path, 'a', encoding='utf-8') as lf:
                lf.write('\n===== %s page %d %s–%s\n' % (now(), n, ids[0], ids[-1])); lf.flush()
                try:
                    proc = subprocess.Popen(cmd, cwd=run_dir(slug), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, env=env)
                except Exception as e:
                    j.update(state='error', message='Could not start Claude Code: %s' % e); return
                j['proc'] = proc; result_text = ''; is_error = False
                for line in proc.stdout:
                    lf.write(line); lf.flush()
                    try: ev = json.loads(line)
                    except ValueError: continue
                    t = ev.get('type')
                    if t == 'assistant':
                        for c in (ev.get('message') or {}).get('content') or []:
                            if c.get('type') == 'text' and c.get('text', '').strip(): j['last'] = c['text'].strip()[-200:]
                            elif c.get('type') == 'tool_use':
                                j['last'] = 'tool: ' + c.get('name', '') + ' ' + str((c.get('input') or {}).get('command') or (c.get('input') or {}).get('file_path') or '')[:120]
                                st = stage_of(c.get('name', ''), c.get('input'))
                                if st is not None and st > (j.get('stage') or 0): j['stage'] = st
                                if st == 3: j['grounding_calls'] = (j.get('grounding_calls') or 0) + 1
                    elif t == 'result':
                        result_text = str(ev.get('result') or ''); is_error = bool(ev.get('is_error'))
                proc.wait()
            if j.get('stop'): j.update(state='stopped', message='Stopped during page %d.' % n); return
            low = result_text.lower()
            if 'not logged in' in low:
                j.update(state='error', message='Claude Code is not logged in for headless runs. Run it once in Terminal and type /login, then start the run again: "%s"' % claude); return
            if is_error and re.search(r'usage limit|rate limit|limit reached|out of extra usage|resets', low):
                t = parse_reset(result_text) or (datetime.datetime.now() + datetime.timedelta(minutes=20))
                j.update(state='waiting', wait_until=t.strftime('%H:%M'), message='Usage limit reached; waiting until %s, then continuing page %d.' % (t.strftime('%H:%M'), n))
                while datetime.datetime.now() < t:
                    if j.get('stop'): j.update(state='stopped', message='Stopped while waiting.'); return
                    time.sleep(15)
                if attempts >= 8: j.update(state='error', message='Gave up on page %d after %d attempts.' % (n, attempts)); return
                continue
            rebuild(slug)
            data = load(os.path.join(pdir(slug), 'data.json'), {})
            still = [u['id'] for u in data.get('units', []) if u.get('pending') and u['id'] in ids]
            if not is_error: record_page_time(slug, ids, (datetime.datetime.now() - j['page_started']).total_seconds(), not still)
            if still and attempts < 3 and not is_error:
                j.update(message='Page %d finished but %s still have no block; retrying (resume).' % (n, ', '.join(still))); continue
            if still: j.update(state='error', message='Page %d: %s have no block in final.md after %d attempts. See the log.' % (n, ', '.join(still), attempts)); return
            break
    j.update(state='done', message='Finished %d page(s).' % len(pages), wait_until=None)

def start_job(slug, scope):
    with JOBLOCK:
        j = JOBS.get(slug)
        if j and j.get('state') in ('running', 'waiting'): return {'ok': False, 'error': 'a run is already going'}
        JOBS[slug] = {'state': 'starting', 'started': now(), 'stop': False, 'message': 'Starting…'}
        th = threading.Thread(target=run_job, args=(slug, scope), daemon=True); th.start()
    return {'ok': True}

def stop_job(slug):
    j = JOBS.get(slug)
    if not j: return {'ok': True}
    j['stop'] = True
    pr = j.get('proc')
    if pr and pr.poll() is None:
        try: pr.terminate()
        except Exception: pass
    return {'ok': True}

def claude_auth_status(slug=None):
    p = pjson(slug) if slug else {}
    c = claude_binary(p or {})
    if not c: return {'found': False, 'logged_in': False, 'message': 'Claude Code was not found on this Mac.'}
    try:
        env = dict(os.environ); env.pop('CLAUDECODE', None)
        r = subprocess.run([c, 'auth', 'status', '--json'], capture_output=True, text=True, timeout=30, env=env)
        j = json.loads(r.stdout.strip() or '{}')
        return {'found': True, 'path': c, 'logged_in': bool(j.get('loggedIn')), 'method': j.get('authMethod'), 'message': 'Logged in (%s).' % j.get('authMethod') if j.get('loggedIn') else 'Not logged in yet.'}
    except Exception as e:
        return {'found': True, 'path': c, 'logged_in': False, 'message': 'Could not read the login status: %s' % e}

def claude_login(slug=None):
    """Open a Terminal window that runs the sign-in; the user finishes it in the browser. Nothing is typed for them."""
    p = pjson(slug) if slug else {}
    c = claude_binary(p or {})
    if not c: return {'ok': False, 'error': 'Claude Code was not found on this Mac.'}
    script = 'tell application "Terminal"\nactivate\ndo script "clear; echo \\"Tiger CAT: signing in to Claude Code. Finish in the browser window that opens, then come back to the app.\\"; " & quoted form of "%s" & " auth login"\nend tell' % c
    try:
        subprocess.run(['osascript', '-e', script], capture_output=True, text=True, timeout=20)
        return {'ok': True, 'message': 'A Terminal window opened with the sign-in. Finish it in your browser; the app will notice.'}
    except Exception as e:
        return {'ok': False, 'error': str(e)}

def check_claude(slug):
    p = pjson(slug) or {}
    c = claude_binary(p)
    if not c: return {'ok': False, 'found': False, 'message': 'No Claude Code binary found.'}
    try:
        env = dict(os.environ); env.pop('CLAUDECODE', None)
        r = subprocess.run([c, '-p', 'Reply with the single word OK.', '--model', 'haiku', '--effort', 'low', '--output-format', 'json'], capture_output=True, text=True, timeout=90, env=env, cwd=pdir(slug))
        try: res = json.loads(r.stdout.strip().splitlines()[-1])
        except Exception: res = {'result': r.stdout[-300:] + r.stderr[-300:], 'is_error': True}
        txt = str(res.get('result', ''))
        if 'not logged in' in txt.lower(): return {'ok': False, 'found': True, 'path': c, 'logged_in': False, 'message': 'Found Claude Code but it is not logged in for headless runs.'}
        if res.get('is_error'): return {'ok': False, 'found': True, 'path': c, 'logged_in': None, 'message': txt[:300]}
        return {'ok': True, 'found': True, 'path': c, 'logged_in': True, 'message': 'Claude Code answers (%s).' % txt.strip()[:40]}
    except Exception as e:
        return {'ok': False, 'found': True, 'path': c, 'message': str(e)}

# ------------------------------------------------------------------ export
def compose_units(slug):
    data = load(os.path.join(pdir(slug), 'data.json'), {})
    state = load(os.path.join(pdir(slug), 'state.json'), {'units': {}}).get('units', {})
    out = []
    for u in data.get('units', []):
        e = state.get(u['id'], {})
        text = e.get('text') or u['text']
        if not text and u.get('pending'): continue
        fns = []
        for f in (e.get('footnotes') if e.get('footnotes') is not None else u['footnotes']):
            if not (f.get('text') or '').strip(): continue
            anchor = (f.get('anchor') or '').strip()
            if anchor == '*': anchor = ''
            if not anchor or not any(anchor in l for l in text): anchor = next((l for l in reversed(text) if l.strip()), '')
            fns.append({'anchor': anchor, 'text': f['text'].strip()})
        out.append({'id': u['id'], 'header': u['header'], 'tib': u['tib'], 'wylie': u['wylie'], 'text': text, 'footnotes': fns, 'notes': u['notes'], 'confidence': u.get('confidence', ''), 'reviewed': e.get('reviewed', False)})
    return data, out

def final_md(data, units):
    L = ['# %s' % data.get('title', 'translation'), '', 'Exported from Tiger CAT on %s; edited text and footnotes from state.json, headers and notes from the run.' % datetime.date.today().isoformat(), '']
    for u in units:
        L += ['### ' + u['id'], 'HEADER: ' + u['header'], 'TEXT:'] + u['text'] + ['FOOTNOTES:']
        L += ['FN(%s): %s' % (f['anchor'], f['text']) for f in u['footnotes']] or ['none']
        L += ['NOTES:'] + (['%s: %s' % (n['tag'], ('(%s) ' % n['qid'] if n.get('qid') else '') + n['text']) for n in u['notes']] or ['none']) + ['']
    return '\n'.join(L)

TIB_FONT = 'Kailasa'          # macOS; Word on Windows falls back to Microsoft Himalaya when Kailasa is absent
EN_FONT = 'Palatino'

def _rfonts(run, font, cs=False):
    """Set the run's font for every script range; Tibetan is a complex script, so w:cs matters most."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(k), font)
    if cs:
        c = OxmlElement('w:cs'); rpr.append(c)

def _tib_run(p, text, size=11):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    r = p.add_run(text); _rfonts(r, TIB_FONT, cs=True); r.font.size = Pt(size)
    rpr = r._element.get_or_add_rPr(); sz = OxmlElement('w:szCs'); sz.set(qn('w:val'), str(size * 2)); rpr.append(sz)
    return r

def _md_runs(p, text, fns=None, pending=None, size=11):
    """Add text with *italics* as italic runs; place pending footnotes after their anchors."""
    from docx.shared import Pt
    pos = []
    if pending is not None:
        for anchor, body in list(pending):
            k = text.find(anchor)
            if k >= 0: pos.append((k + len(anchor), anchor, body)); pending.remove([anchor, body])
        pos.sort()
    cuts = [0] + [k for k, _, _ in pos] + [len(text)]
    for i in range(len(cuts) - 1):
        seg = text[cuts[i]:cuts[i + 1]]
        for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', seg):
            if not part: continue
            bold = part.startswith('**'); ital = not bold and part.startswith('*')
            r = p.add_run(part.strip('*') if (bold or ital) else part); _rfonts(r, EN_FONT); r.font.size = Pt(size); r.italic = ital; r.bold = bold
        if i < len(pos): fns.reference(p, _fn_add(fns, pos[i][2]))

def _fn_add(fns, text):
    """A footnote whose *italics* (titles) come out as italic runs; same XML as the skill's Footnotes.add."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    fid = fns.next_id; fns.next_id += 1
    fn = OxmlElement('w:footnote'); fn.set(qn('w:id'), str(fid))
    p = OxmlElement('w:p'); ppr = OxmlElement('w:pPr'); sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '0'); ppr.append(sp); p.append(ppr)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); va = OxmlElement('w:vertAlign'); va.set(qn('w:val'), 'superscript'); rpr.append(va); r.append(rpr); r.append(OxmlElement('w:footnoteRef')); p.append(r)
    first = True
    for part in re.split(r'(\*[^*]+\*)', text):
        if not part: continue
        ital = part.startswith('*') and part.endswith('*') and len(part) > 2
        r2 = OxmlElement('w:r'); rpr2 = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '20'); rpr2.append(sz)
        if ital: rpr2.append(OxmlElement('w:i'))
        r2.append(rpr2); t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = (' ' if first else '') + (part.strip('*') if ital else part); r2.append(t); p.append(r2); first = False
    fn.append(p); fns.root.append(fn)
    return fid

def _fill_cell(cell, u, fns, size=11):
    first = True; paras, buf = [], []
    for l in u['text'] + ['']:
        if l.strip(): buf.append(l)
        elif buf: paras.append(buf); buf = []
    pending = [[f['anchor'], f['text']] for f in u['footnotes']]
    for lines in paras:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        if len(lines) == 1: _md_runs(p, lines[0], fns, pending, size)
        else:
            for i, l in enumerate(lines):
                if i: p.add_run().add_break()
                _md_runs(p, l, fns, pending, size)
    for anchor, body in pending: fns.reference(cell.paragraphs[-1], _fn_add(fns, body))

def build_docx(units, out, title, notes, aligned, headers=False):
    sys.path.insert(0, SKILL_TOOLS)
    import export_docx as X, docx
    from docx.shared import Pt, Cm
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = docx.Document(); fns = X.Footnotes(doc)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(2.0)
    st = doc.styles['Normal']; st.font.name = EN_FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:cs'), TIB_FONT)
    if title:
        h = doc.add_heading(level=1); r = h.add_run(re.sub(r'\*', '', title)); _rfonts(r, EN_FONT)
    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    if aligned:
        table = doc.add_table(rows=0, cols=2); table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        tblPr = table._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        w0, w1 = int(usable * 0.46), int(usable * 0.54)          # EMU for python-docx; the grid wants twips
        grid = table._tbl.find(qn('w:tblGrid'))
        for gc, w in zip(grid.findall(qn('w:gridCol')), (w0, w1)): gc.set(qn('w:w'), str(int(w / 635)))
        for u in units:
            cells = table.add_row().cells
            cells[0].width = w0; cells[1].width = w1
            p = cells[0].paragraphs[0]
            r = p.add_run(u['id'] + ('  ·  ' + u['confidence'] if headers and u.get('confidence') else '') + '\n'); _rfonts(r, EN_FONT); r.font.size = Pt(7); r.italic = True
            _tib_run(p, u['tib'], 11)
            pf = p.paragraph_format; pf.line_spacing = 1.5
            _fill_cell(cells[1], u, fns, 11)
            if headers and u.get('confidence') in ('low', 'very low'):
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'FDE4CF' if u['confidence'] == 'low' else 'F8CFCB'); cells[1]._tc.get_or_add_tcPr().append(shd)
    else:
        for u in units:
            if headers:
                h = doc.add_paragraph(); r = h.add_run(u['id'] + (' · ' + u['confidence'] if u.get('confidence') else '')); _rfonts(r, EN_FONT); r.italic = True; r.font.size = Pt(8)
            paras, buf = [], []
            for l in u['text'] + ['']:
                if l.strip(): buf.append(l)
                elif buf: paras.append(buf); buf = []
            pending = [[f['anchor'], f['text']] for f in u['footnotes']]
            for lines in paras:
                p = doc.add_paragraph()
                if len(lines) == 1: _md_runs(p, lines[0], fns, pending, 11)
                else:
                    for i, l in enumerate(lines):
                        if i: p.add_run().add_break()
                        _md_runs(p, l, fns, pending, 11)
            for anchor, body in pending: fns.reference(doc.paragraphs[-1], _fn_add(fns, body))
    if notes and any(u['notes'] for u in units):
        doc.add_page_break(); doc.add_heading("Editor's notes", level=1)
        for u in units:
            if u['notes']:
                doc.add_paragraph(u['id'] + (' · confidence: ' + u['confidence'] if u['confidence'] else ''), style='Heading 3')
                for n in u['notes']:
                    p = doc.add_paragraph(); _md_runs(p, '%s: %s' % (n['tag'], n['text']), size=9)
    fns.save_into(); doc.save(out)

def do_export(slug, body):
    mode = body.get('mode', 'en'); notes = bool(body.get('notes')); headers = bool(body.get('headers'))
    data, units = compose_units(slug)
    if not units: return {'ok': False, 'error': 'nothing translated yet'}
    exp = os.path.join(pdir(slug), 'export'); os.makedirs(exp, exist_ok=True)
    stamp = datetime.datetime.now().strftime('%Y%m%d-%H%M'); md_path = os.path.join(exp, 'final.edited.md')
    open(md_path, 'w', encoding='utf-8').write(final_md(data, units))
    if mode == 'md': return {'ok': True, 'file': '/export/final.edited.md?p=' + slug, 'name': 'final.edited.md', 'log': '%d units' % len(units)}
    title = data.get('title', ''); nf = sum(len(u['footnotes']) for u in units)
    name = '%s-%s-%s.docx' % (slug, 'aligned' if mode == 'aligned' else 'en', stamp)
    build_docx(units, os.path.join(exp, name), title, notes, aligned=(mode == 'aligned'), headers=headers)
    return {'ok': True, 'file': '/export/%s?p=%s' % (name, slug), 'name': name, 'log': '%d units, %d footnotes' % (len(units), nf)}

# ------------------------------------------------------------------ http
class H(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *a, **k): super().__init__(*a, directory=APP, **k)
    def log_message(self, fmt, *args):
        line = fmt % args
        if '/api/refresh' in line or '/api/run?' in line or '/api/usage' in line or ('" 200 ' in line and '/api/' not in line): return
        sys.stderr.write(line + '\n')
    def _json(self, obj, code=200):
        b = json.dumps(obj, ensure_ascii=False).encode('utf-8')
        self.send_response(code); self.send_header('Content-Type', 'application/json; charset=utf-8'); self.send_header('Content-Length', str(len(b))); self.end_headers(); self.wfile.write(b)
    def q(self):
        u = urllib.parse.urlparse(self.path); return u.path, dict(urllib.parse.parse_qsl(u.query))
    def translate_path(self, path):
        p, qs = self.q(); slug = qs.get('p', '')
        if p == '/data.json' and slug: return os.path.join(pdir(slug), 'data.json')
        if p.startswith('/export/') and slug: return os.path.join(pdir(slug), 'export', os.path.basename(p))
        return super().translate_path(path)
    def end_headers(self):
        if self.path.startswith('/export/'): self.send_header('Content-Disposition', 'attachment')
        self.send_header('Cache-Control', 'no-store'); super().end_headers()
    def do_GET(self):
        p, qs = self.q(); slug = qs.get('p', '')
        try:
            if p == '/api/projects': return self._json({'projects': list_projects(), 'claude': claude_binary(), 'audiences': AUDIENCES, 'purposes': PURPOSES, 'notes': NOTES})
            if p == '/api/project' and slug:
                pj = pjson(slug)
                if not pj: return self._json({'ok': False, 'error': 'no project'}, 404)
                units_text = ''
                if os.path.exists(units_path(slug)): units_text = '\n'.join(l for l in open(units_path(slug), encoding='utf-8').read().splitlines() if re.match(r'^U\d+', l))
                return self._json({'ok': True, 'project': pj, 'units_text': units_text, 'run_dir': run_dir(slug), 'claude': claude_binary(pj)})
            if p == '/api/state' and slug: return self._json(load(os.path.join(pdir(slug), 'state.json'), {'units': {}}))
            if p == '/api/refresh' and slug: return self._json(api_refresh(slug))
            if p == '/api/run' and slug: return self._json(job_status(slug))
            if p == '/api/run/log' and slug:
                j = JOBS.get(slug); lp = (j or {}).get('log')
                if not lp: logs = sorted(glob.glob(os.path.join(pdir(slug), 'jobs', '*.log'))); lp = logs[-1] if logs else None
                if not lp or not os.path.exists(lp): return self._json({'log': ''})
                txt = open(lp, encoding='utf-8', errors='replace').read()[-20000:]
                return self._json({'log': txt, 'file': lp})
            if p == '/api/claude/check' and slug: return self._json(check_claude(slug))
            if p == '/api/claude/status': return self._json(claude_auth_status(slug or None))
            if p == '/api/usage': return self._json(api_usage())
        except Exception as e:
            traceback.print_exc(); return self._json({'ok': False, 'error': str(e)}, 500)
        return super().do_GET()
    def do_POST(self):
        p, qs = self.q(); slug = qs.get('p', '')
        n = int(self.headers.get('Content-Length', 0) or 0)
        try: body = json.loads(self.rfile.read(n) or b'{}')
        except Exception: return self._json({'ok': False, 'error': 'bad json'}, 400)
        try:
            if p == '/api/projects': return self._json(create_project(body))
            if p == '/api/project' and slug: return self._json(update_project(slug, body))
            if p == '/api/project/delete' and slug: return self._json(delete_project(slug))
            if p == '/api/project/resegment' and slug: return self._json(resegment_project(slug, body))
            if p == '/api/state' and slug: dump(os.path.join(pdir(slug), 'state.json'), body); return self._json({'ok': True})
            if p == '/api/export' and slug: return self._json(do_export(slug, body))
            if p == '/api/glossary/import' and slug: return self._json(import_glossary(slug, body.get('text') or ''))
            if p == '/api/run' and slug: return self._json(start_job(slug, body.get('scope') or 'next'))
            if p == '/api/run/stop' and slug: return self._json(stop_job(slug))
            if p == '/api/claude/login': return self._json(claude_login(slug or None))
        except Exception as e:
            traceback.print_exc(); return self._json({'ok': False, 'error': str(e)}, 500)
        self._json({'ok': False, 'error': 'not found'}, 404)

def main():
    global PROJECTS
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--port', type=int, default=8765); ap.add_argument('--projects', default=PROJECTS)
    a = ap.parse_args(); PROJECTS = os.path.abspath(a.projects); os.makedirs(PROJECTS, exist_ok=True)
    srv = http.server.ThreadingHTTPServer(('127.0.0.1', a.port), H)
    inst, repo = os.path.join(INSTALLED_SKILL, 'SKILL.md'), os.path.join(REPO_SKILL, 'SKILL.md')
    if not os.path.isfile(inst): print('skills are not installed in ~/.claude/skills; run ./install.sh --skills-only (runs need the installed copy)', flush=True)
    elif os.path.isfile(repo) and os.path.getmtime(inst) < os.path.getmtime(repo): print('skills in ~/.claude/skills are older than the repo; run ./install.sh --skills-only', flush=True)
    print('Tiger CAT at http://127.0.0.1:%d  (projects in %s; Claude Code: %s)' % (a.port, PROJECTS, claude_binary() or 'not found'), flush=True)
    try: srv.serve_forever()
    except KeyboardInterrupt: pass

if __name__ == '__main__':
    main()
